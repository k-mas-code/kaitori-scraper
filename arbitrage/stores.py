"""店舗リスト (data/yahoo_pp_stores_5genres.xlsx) の読み込み"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

from .config import ROOT

DEFAULT_STORES_PATH = ROOT / "data" / "yahoo_pp_stores_5genres.xlsx"

SHEET_ALL = "全ジャンル"
SHEET_UPSELL = "ポイント上乗せ調査"
SHEET_EXCLUDED = "除外リスト（非適格）"
# 上乗せ調査シートに無い店の最大上乗せ率 (絞り込み用の上限値)
DEFAULT_MAX_UPSELL = Decimal("0.14")

# 日別列のヘッダーは「10/4\n日」の形 (「10/4以降も掲載予定」等の列と区別する)
DATE_HEADER_RE = re.compile(r"^\s*(\d{1,2})/(\d{1,2})\s*[月火水木金土日]\s*$")


class StoreListError(ValueError):
    pass


@dataclass(frozen=True)
class Store:
    store_id: str
    name: str
    genre: str
    good_store: bool
    bsplus_rate: Decimal          # 実行日のボーナスストアPlus枠 (0 / 0.04 / 0.09)
    max_upsell: Decimal           # ストアポイント上乗せ率の上限 (絞り込み用)


@dataclass(frozen=True)
class StoreList:
    stores: dict[str, Store]      # store_id → Store (除外リストの店は含まない)
    run_date: date
    global_bonus_labels: str | None   # 実行日の全体加算の表記 ("+2" / "+2/+3" / None)
    source_file: str


def _header_index(header_row, title: str) -> int:
    for i, value in enumerate(header_row):
        if value is not None and str(value).replace("\n", "").strip() == title:
            return i
    raise StoreListError(f"列「{title}」が見つからない")


def _date_column(header_row, run_date: date) -> int:
    for i, value in enumerate(header_row):
        m = DATE_HEADER_RE.match(str(value)) if value is not None else None
        if m and (int(m.group(1)), int(m.group(2))) == (run_date.month, run_date.day):
            return i
    raise StoreListError(f"店舗リストに {run_date:%m/%d} の列が無い")


def load_stores(run_date: date, path: Path = DEFAULT_STORES_PATH) -> StoreList:
    if not path.exists():
        raise StoreListError(f"{path} が無い")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for sheet in (SHEET_ALL, SHEET_UPSELL, SHEET_EXCLUDED):
        if sheet not in wb.sheetnames:
            raise StoreListError(f"シート「{sheet}」が無い")

    # 除外リスト: 1 行目=説明, 2 行目=ヘッダー
    rows = list(wb[SHEET_EXCLUDED].iter_rows(values_only=True))
    col = _header_index(rows[1], "ストアID")
    excluded = {str(r[col]).strip() for r in rows[2:] if r[col]}

    # 上乗せ調査: 4 行目=ヘッダー
    rows = list(wb[SHEET_UPSELL].iter_rows(values_only=True))
    col_id, col_max = _header_index(rows[3], "ストアID"), _header_index(rows[3], "最大上乗せ率")
    col_unknown = _header_index(rows[3], "判定できなかった商品数")
    max_upsell: dict[str, Decimal] = {}
    for r in rows[4:]:
        if not r[col_id]:
            continue
        surveyed = Decimal(str(r[col_max] or 0))
        if not 0 <= surveyed <= 1:
            raise StoreListError(f"{r[col_id]}: 最大上乗せ率が想定外の値 {r[col_max]!r} (0.14 = 14%)")
        # 判定できなかった商品がある店は、調査値より高い上乗せがありうるので既定値を下回らせない
        max_upsell[str(r[col_id]).strip()] = max(surveyed, DEFAULT_MAX_UPSELL) if r[col_unknown] else surveyed

    # 全ジャンル: 1 行目=全体加算/休み, 2 行目=ヘッダー
    rows = list(wb[SHEET_ALL].iter_rows(values_only=True))
    flags, header = rows[0], rows[1]
    col_id = _header_index(header, "ストアID")
    col_name = _header_index(header, "ストア名")
    col_genre = _header_index(header, "ジャンル")
    col_good = _header_index(header, "優良ストア")
    # 日別列のヘッダーは月日だけなので、掲載期間 (初回/最終掲載日) で年を確かめる
    col_first, col_last = _header_index(header, "初回掲載日"), _header_index(header, "最終掲載日")
    firsts = [r[col_first].date() for r in rows[2:] if isinstance(r[col_first], datetime)]
    lasts = [r[col_last].date() for r in rows[2:] if isinstance(r[col_last], datetime)]
    if not firsts or not lasts or not min(firsts) <= run_date <= max(lasts):
        period = f"{min(firsts)}〜{max(lasts)}" if firsts and lasts else "不明"
        raise StoreListError(
            f"{run_date} は店舗リストの掲載期間 ({period}) の外。新しいリストに差し替える")
    col_date = _date_column(header, run_date)
    flag = str(flags[col_date]).strip() if col_date < len(flags) and flags[col_date] else None
    global_bonus_labels = flag if flag and flag.startswith("+") else None

    stores: dict[str, Store] = {}
    for r in rows[2:]:
        if not r[col_id]:
            continue
        store_id = str(r[col_id]).strip()
        if store_id in excluded:
            continue
        slot = r[col_date] if col_date < len(r) else None
        if slot in (None, ""):
            bsplus = Decimal(0)
        elif slot in (4, 9):
            bsplus = Decimal(int(slot)) / 100
        else:
            raise StoreListError(f"{store_id}: {run_date:%m/%d} の枠が想定外の値 {slot!r}")
        stores[store_id] = Store(
            store_id=store_id,
            name=str(r[col_name] or ""),
            genre=str(r[col_genre] or ""),
            good_store=str(r[col_good] or "").strip() == "○",
            bsplus_rate=bsplus,
            max_upsell=max_upsell.get(store_id, DEFAULT_MAX_UPSELL),
        )
    if not stores:
        raise StoreListError("店舗が 1 件も読めなかった")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()[:16]
    return StoreList(stores=stores, run_date=run_date, global_bonus_labels=global_bonus_labels,
                     source_file=f"{path.name} (sha256:{digest})")
